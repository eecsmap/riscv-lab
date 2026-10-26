#!/usr/bin/env python3
"""Build the step-3 CSR lecture deck (Chinese) as .pptx, for import into Google Slides.

    python3 build_csr_deck.py [out.pptx]          # needs python-pptx

Latin text is Arial, Chinese text Noto Sans SC (both available in Google Slides), code Courier New.
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

OUT = sys.argv[1] if len(sys.argv) > 1 else "CSR.pptx"

# ---- palette: deep teal dominates, amber is the one sharp accent
NAVY   = RGBColor(0x0F, 0x2E, 0x3D)   # dark slides, headings
TEAL   = RGBColor(0x2A, 0x7F, 0x86)   # secondary
TINT   = RGBColor(0xE9, 0xF2, 0xF3)   # card background on white
AMBER  = RGBColor(0xE8, 0xA3, 0x3D)   # accent
INK    = RGBColor(0x1B, 0x2B, 0x33)   # body text
MUTED  = RGBColor(0x5B, 0x6B, 0x73)   # captions
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GRAY   = RGBColor(0xC9, 0xD3, 0xD6)
LATIN, CJK, MONO = "Arial", "Noto Sans SC", "Courier New"

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # exact 16:9, no sliver at the edge
BLANK = prs.slide_layouts[6]


def _fonts(run, mono=False):
    """latin + east-asian typefaces on one run, in schema order (after solidFill)."""
    run.font.name = MONO if mono else LATIN
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        old = rpr.find(qn(tag))
        if old is not None:
            rpr.remove(old)
    latin = rpr.find(qn("a:latin"))
    ea = etree.SubElement(rpr, qn("a:ea")); ea.set("typeface", MONO if mono else CJK)
    latin.addnext(ea)


def text(slide, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, mono=False, space_after=6, margin=0.05):
    """paras: list of str or list of (str, {opts}) runs per paragraph."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    first = True
    for para in paras:
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = align; p.space_after = Pt(space_after)
        runs = para if isinstance(para, list) else [(para, {})]
        for s, o in runs:
            r = p.add_run(); r.text = s
            r.font.size = Pt(o.get("size", size)); r.font.bold = o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
            _fonts(r, o.get("mono", mono))
    return tb


def box(slide, x, y, w, h, fill, radius=0.08, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(1.25)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def chip(slide, x, y, addr, name, w=2.0, dark=True):
    """the deck's motif: a register chip, hex address over its name."""
    box(slide, x, y, w, 0.82, NAVY if dark else WHITE, radius=0.18, line=None if dark else TEAL)
    text(slide, x, y + 0.06, w, 0.36, [addr], size=14, color=AMBER, mono=True, bold=True, align=PP_ALIGN.CENTER)
    text(slide, x, y + 0.40, w, 0.36, [name], size=14, color=WHITE if dark else NAVY, mono=True, align=PP_ALIGN.CENTER)


def title(slide, s, dark=False):
    text(slide, 0.6, 0.45, 12.1, 0.8, [s], size=34, bold=True, color=WHITE if dark else NAVY)


def bg(slide, color):
    f = slide.background.fill; f.solid(); f.fore_color.rgb = color


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def circle_num(slide, x, y, n, d=0.6):
    c = box(slide, x, y, d, d, AMBER, shape=MSO_SHAPE.OVAL)
    text(slide, x, y, d, d, [str(n)], size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)


# ============================================================ 1 title (dark)
s = prs.slides.add_slide(BLANK); bg(s, NAVY)
text(s, 0.8, 1.55, 11.7, 1.2, ["CSR：核自己的控制面板"], size=48, bold=True, color=WHITE)
text(s, 0.8, 2.75, 11.7, 0.6, ["RISC-V 控制和状态寄存器（Control and Status Register）"], size=22, color=GRAY)
text(s, 0.8, 3.35, 11.7, 0.5, ["多周期 RV64 教学路线 · 第 3 步讲义"], size=16, color=AMBER)
for i, (a, n) in enumerate([("0x300", "mstatus"), ("0x305", "mtvec"), ("0x340", "mscratch"),
                            ("0x341", "mepc"), ("0x342", "mcause"), ("0x343", "mtval")]):
    chip(s, 0.8 + i * 1.98, 5.15, a, n, w=1.78, dark=False)
notes(s, "这一讲回答三个问题：CSR 是什么、为什么不能用普通内存代替、第 3 步的核为什么必须实现这 12 个地址。"
         "底下这一排就是 trap 机制要用到的 6 个 CSR，后面会逐个讲。")

# ============================================================ 2 the third address space
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
title(s, "和寄存器、内存并列的第三个地址空间")
cols = [("通用寄存器", ["32 个", "几乎所有指令", "程序的数据"], False),
        ("CSR", ["地址 12 位，最多 4096 个", "只有 6 条 Zicsr 指令", "核自己的配置和状态"], True),
        ("内存", ["整个物理地址空间", "load / store", "程序、数据、外设"], False)]
labels = ["数量", "怎么访问", "装什么"]
for i, (head, vals, hi) in enumerate(cols):
    x = 0.6 + i * 4.1
    box(s, x, 1.55, 3.8, 4.35, NAVY if hi else TINT)
    text(s, x + 0.3, 1.75, 3.2, 0.6, [head], size=24, bold=True, color=AMBER if hi else NAVY)
    for j, (lab, v) in enumerate(zip(labels, vals)):
        yy = 2.6 + j * 1.05
        text(s, x + 0.3, yy, 3.2, 0.35, [lab], size=13, color=GRAY if hi else MUTED)
        text(s, x + 0.3, yy + 0.33, 3.3, 0.6, [v], size=17, bold=True, color=WHITE if hi else INK)
text(s, 0.6, 6.2, 12.1, 0.6, ["普通程序几乎不碰 CSR：它们是给操作系统和固件用的，用来配置核、记录刚才发生了什么、测量性能。"],
     size=16, color=MUTED)
notes(s, "CSR 是 RISC-V 里第三个地址空间。通用寄存器装程序的数据，内存装程序本身和外设，CSR 装的是核自己的状态。"
         "访问它只有 csrrw、csrrs、csrrc 和三个立即数形式，一共 6 条指令。")

# ============================================================ 3 the address encodes the rules
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
title(s, "地址本身编码了访问规则")
bx, by, cw = 0.95, 1.6, 0.9
groups = [(11, 10, AMBER, "可写性", "11 = 只读"), (9, 8, TEAL, "最低特权级", "00 U · 01 S · 11 M"),
          (7, 0, GRAY, "编号", "同一类里的序号")]
for bit in range(12):
    col = AMBER if bit >= 10 else TEAL if bit >= 8 else GRAY
    xx = bx + (11 - bit) * (cw + 0.06)
    box(s, xx, by, cw, 0.7, col, radius=0.12)
    text(s, xx, by, cw, 0.7, [str(bit)], size=15, bold=True, color=NAVY if col != TEAL else WHITE,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)
for hi, lo, col, name, rule in groups:
    x0 = bx + (11 - hi) * (cw + 0.06); x1 = bx + (11 - lo) * (cw + 0.06) + cw
    text(s, x0 - 0.2, by + 0.82, x1 - x0 + 0.4, 0.4, [f"[{hi}:{lo}] {name}"], size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    text(s, x0 - 0.2, by + 1.22, x1 - x0 + 0.4, 0.4, [rule], size=14, color=MUTED, align=PP_ALIGN.CENTER)
examples = [("0xC00", "cycle", "11 · 00", "只读，U 模式可读"), ("0x300", "mstatus", "00 · 11", "可写，只有 M 模式"),
            ("0xF14", "mhartid", "11 · 11", "只读，只有 M 模式"), ("0xB00", "mcycle", "10 · 11", "可写，只有 M 模式")]
for i, (a, n, bits, mean) in enumerate(examples):
    x = 0.6 + i * 3.1
    chip(s, x, 3.75, a, n, w=2.8)
    text(s, x, 4.67, 2.8, 0.4, [bits], size=15, mono=True, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    text(s, x, 5.05, 2.8, 0.4, [mean], size=15, color=INK, align=PP_ALIGN.CENTER)
text(s, 0.6, 6.15, 12.1, 0.55, [[("所以只读判断可以直接写成  ", {}), ("read_only = (addr[11:10] == 2'b11);", {"mono": True, "bold": True, "color": NAVY})]],
     size=16, color=MUTED)
notes(s, "12 位地址的最高两位决定可不可写，11 就是只读；接下来两位是访问它需要的最低特权级。"
         "第 6 步加特权级检查时，用的就是 addr[9:8]。")

# ============================================================ 4 why not MMIO
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
title(s, "为什么不放进内存地址空间")
reasons = [("处理程序要在动内存之前就能用", "trap 刚发生时连栈指针都不可信，第一件事是", "csrrw sp, mscratch, sp"),
           ("读改写必须原子", "一条指令读出旧值、置位、写回，中间不能被打断", "csrrs t0, mstatus, t1"),
           ("每个 hart 一份", "多核时每个核有自己的 mepc、mcause，不能共用一个内存地址", None),
           ("改了要立即生效", "写 mstatus.MIE，下一条指令就按新开关执行；写 satp，下一次取指就按新页表翻译", None)]
for i, (h, d, code) in enumerate(reasons):
    x = 0.6 + (i % 2) * 6.15; y = 1.55 + (i // 2) * 2.55
    box(s, x, y, 5.95, 2.3, TINT)
    circle_num(s, x + 0.3, y + 0.3, i + 1)
    text(s, x + 1.1, y + 0.3, 4.6, 0.55, [h], size=20, bold=True, color=NAVY)
    text(s, x + 1.1, y + 0.9, 4.65, 0.8, [d], size=15, color=INK)
    if code:
        text(s, x + 1.1, y + 1.7, 4.6, 0.4, [code], size=14, mono=True, bold=True, color=TEAL)
notes(s, "CLINT 和 GPIO 都是内存地址，为什么 CSR 不是？四个原因。最根本的是第一条："
         "处理程序要先从 mscratch 里换出一个可信的栈，才能开始用内存，如果 mscratch 本身在内存里，就成了鸡生蛋。")

# ============================================================ 5 the twelve of step 3
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
title(s, "第 3 步实现的 12 个地址")
box(s, 0.6, 1.6, 3.2, 4.9, NAVY)
text(s, 0.6, 2.05, 3.2, 1.6, ["12"], size=96, bold=True, color=AMBER, align=PP_ALIGN.CENTER)
text(s, 0.75, 3.85, 2.9, 1.5, ["个地址", "其中 8 个真正存东西，其余是常量或别名"], size=16, color=WHITE, align=PP_ALIGN.CENTER)
grp = [("trap 机制 · 6 个", [("0x305", "mtvec"), ("0x341", "mepc"), ("0x342", "mcause"), ("0x343", "mtval"), ("0x340", "mscratch"), ("0x300", "mstatus")]),
       ("身份 · 2 个", [("0x301", "misa"), ("0xF14", "mhartid")]),
       ("计数器 · 4 个地址，2 个寄存器", [("0xB00", "mcycle"), ("0xB02", "minstret"), ("0xC00", "cycle"), ("0xC02", "instret")])]
y = 1.6
for g, items in grp:
    text(s, 4.2, y, 8.5, 0.45, [g], size=18, bold=True, color=NAVY)
    for i, (a, n) in enumerate(items):
        chip(s, 4.2 + i * 1.45, y + 0.5, a, n, w=1.35, dark=False)
    y += 1.65
notes(s, "按用途分三组。trap 机制的 6 个是这一步的核心；身份的两个是常量；计数器有 4 个地址，"
         "但 cycle、instret 只是 mcycle、minstret 的只读别名。")

# ============================================================ 6 the six trap CSRs
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
title(s, "trap 机制的 6 个 CSR")
rows = [("CSR", "作用", "谁写", "谁读"),
        ("mtvec", "trap 去哪里：处理程序的地址", "软件启动时", "硬件 trap 时作为新 pc"),
        ("mepc", "trap 发生在哪：出错指令的 pc", "硬件 trap 时", "mret；处理程序可以改它"),
        ("mcause", "为什么：cause 编号，最高位区分中断", "硬件", "处理程序据此分发"),
        ("mtval", "附加信息：出错的地址或指令编码", "硬件", "处理程序"),
        ("mscratch", "处理程序的暂存，通常放它的栈指针", "软件", "软件"),
        ("mstatus", "MIE 中断总开关；MPIE、MPP 保存 trap 前的状态", "软件；trap 和 mret 自动更新", "硬件和软件")]
widths = [1.7, 5.0, 2.7, 2.7]
tbl = s.shapes.add_table(len(rows), 4, Inches(0.6), Inches(1.55), Inches(sum(widths)), Inches(4.9)).table
for j, wv in enumerate(widths):
    tbl.columns[j].width = Inches(wv)
for i, row in enumerate(rows):
    tbl.rows[i].height = Inches(0.66)
    for j, v in enumerate(row):
        c = tbl.cell(i, j); c.margin_left = c.margin_right = Inches(0.12)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.fill.solid(); c.fill.fore_color.rgb = NAVY if i == 0 else (TINT if i % 2 else WHITE)
        tf = c.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]; p.text = ""
        r = p.add_run(); r.text = v
        r.font.size = Pt(15); r.font.bold = (i == 0 or j == 0)
        r.font.color.rgb = WHITE if i == 0 else (TEAL if j == 0 else INK)
        _fonts(r, mono=(j == 0 and i > 0))
notes(s, "这张表的关键是“谁写谁读”。mepc、mcause、mtval 由硬件在 trap 时写，处理程序读；"
         "mtvec 和 mscratch 由软件准备好；mstatus 两边都动：软件开关中断，硬件在 trap 和 mret 时自动保存和恢复。")

# ============================================================ 7 one trap, step by step
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
title(s, "一次 trap 的经过")
steps = [("出错", ["非法指令", "未对齐访问", "访问错误", "ecall、ebreak"], False, "指令"),
         ("保存现场", ["mepc ← pc", "mcause ← 原因", "mtval ← 地址", "MPIE ← MIE", "MIE  ← 0"], True, "硬件"),
         ("跳转", ["pc ← mtvec"], True, "硬件"),
         ("处理", ["csrrw sp,", "  mscratch, sp", "读 mcause 分发", "mepc ← mepc+4"], True, "软件"),
         ("返回", ["mret:", "pc  ← mepc", "MIE ← MPIE"], True, "软件")]
bw, gap = 2.15, 0.36
for i, (h, lines, code, who) in enumerate(steps):
    x = 0.6 + i * (bw + gap)
    box(s, x, 1.65, bw, 3.45, NAVY if who == "硬件" else TINT)
    circle_num(s, x + 0.2, 1.85, i + 1, d=0.55)
    text(s, x + 0.85, 1.87, bw - 0.9, 0.5, [h], size=19, bold=True, color=WHITE if who == "硬件" else NAVY)
    text(s, x + 0.2, 2.5, bw - 0.3, 0.35, [who], size=13, color=AMBER if who == "硬件" else MUTED)
    text(s, x + 0.15, 3.0, bw - 0.25, 2.6, lines, size=13 if code else 15, mono=code,
         color=WHITE if who == "硬件" else INK, space_after=4)
    if i < len(steps) - 1:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + bw + 0.05), Inches(3.2), Inches(gap - 0.1), Inches(0.4))
        a.fill.solid(); a.fill.fore_color.rgb = AMBER; a.line.fill.background()
text(s, 0.6, 5.45, 12.1, 0.6, [[("tests/cpu/tests2/thandler.S", {"mono": True, "bold": True, "color": NAVY}),
                                ("  就是一个只用这 6 个 CSR 的最小处理程序。", {})]], size=16, color=MUTED)
notes(s, "深色的两步由硬件完成，在 S_TRAP 那一拍一次做完；浅色的由软件完成。"
         "mepc 加 4 是为了跳过出错的那条指令；如果不改，mret 会回到同一条指令，再出错一次。")

# ============================================================ 8 identity and counters
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
title(s, "身份与计数器")
box(s, 0.6, 1.55, 5.95, 4.95, TINT)
text(s, 0.9, 1.75, 5.4, 0.5, ["身份：misa、mhartid"], size=20, bold=True, color=NAVY)
text(s, 0.9, 2.35, 5.4, 0.6, ["0x8000_0000_0000_0100"], size=24, bold=True, mono=True, color=TEAL)
text(s, 0.9, 3.0, 5.4, 0.9, ["最高两位 MXL = 2：RV64", "bit 8：I，只有基本整数指令"], size=15, color=INK, space_after=2)
text(s, 0.9, 3.85, 5.4, 1.2, ["只报真正实现的扩展：多报一个，软件就会去用它然后崩溃。c02 的 code 5 测的就是这一条。"], size=15, color=INK)
text(s, 0.9, 4.85, 5.4, 1.0, ["mhartid：我是第几个核。xv6 启动时 0 号核做初始化，其他核等待。"], size=15, color=INK)
box(s, 6.8, 1.55, 5.95, 4.95, NAVY)
text(s, 7.1, 1.75, 5.4, 0.5, ["计数器：mcycle、minstret"], size=20, bold=True, color=AMBER)
text(s, 7.1, 2.35, 5.4, 0.6, ["CPI = Δmcycle ÷ Δminstret"], size=22, bold=True, mono=True, color=WHITE)
text(s, 7.1, 3.1, 5.4, 1.6, ["mcycle 每拍加一", "minstret 每退休一条指令加一，trap 不算", "cycle、instret 是给低特权级读的只读别名"],
     size=15, color=WHITE, space_after=5)
text(s, 7.1, 4.75, 5.4, 1.5, [[("这个项目的所有性能数字都来自它们。基线核在板上跑 32 位 ALU 循环：", {}),
                               ("CPI 40.3", {"bold": True, "color": AMBER})]], size=15, color=GRAY)
notes(s, "misa 的值必须诚实，它是软件判断能不能用某条指令的唯一依据。"
         "计数器方面要记住 trap 不算退休：c10 用 minstret 的差值恰好是 4 来检查这一点。")

# ============================================================ 9 why only now
s = prs.slides.add_slide(BLANK); bg(s, WHITE)
title(s, "为什么到第 3 步才需要它们")
text(s, 0.6, 1.45, 12.1, 0.9, ["第 1、2 步的核只会顺序执行。第 3 步要求核在出错时停下、记录现场、跳到处理程序，还能回来。"],
     size=17, color=INK)
pairs = [("出错后去哪里", ["mtvec"]), ("现场是什么", ["mepc", "mcause", "mtval"]),
         ("处理程序怎么开始工作", ["mscratch"]), ("回来时恢复什么", ["mstatus"])]
for i, (need, csrs) in enumerate(pairs):
    y = 2.45 + i * 0.98
    box(s, 0.6, y, 4.6, 0.78, TINT)
    text(s, 0.85, y, 4.2, 0.78, [need], size=18, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(5.4), Inches(y + 0.2), Inches(0.7), Inches(0.38))
    a.fill.solid(); a.fill.fore_color.rgb = AMBER; a.line.fill.background()
    for k, n in enumerate(csrs):
        box(s, 6.3 + k * 2.1, y, 1.95, 0.78, NAVY, radius=0.18)
        text(s, 6.3 + k * 2.1, y, 1.95, 0.78, [n], size=16, bold=True, mono=True, color=WHITE,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=0)
notes(s, "没有这几个 CSR，trap 就只是一次跳转：回不来，也不知道为什么。每个需求正好对应一个或几个 CSR。")

# ============================================================ 10 what comes next (dark)
s = prs.slides.add_slide(BLANK); bg(s, NAVY)
title(s, "后面的步骤还会加", dark=True)
tl = [("第 3 步", "12 个地址", "trap、身份、计数器"), ("第 4 步", "mie、mip", "中断分线开关与挂起位"),
      ("第 6 步", "S 模式一整套", "sstatus、sepc、stvec……  委托 medeleg、mideleg"), ("第 7 步", "satp", "页表根地址，开启 Sv39")]
s.shapes.add_connector(1, Inches(0.9), Inches(3.05), Inches(12.4), Inches(3.05)).line.color.rgb = TEAL
for i, (st, what, det) in enumerate(tl):
    x = 0.9 + i * 3.1
    box(s, x, 2.78, 0.55, 0.55, AMBER, shape=MSO_SHAPE.OVAL)
    text(s, x, 1.95, 2.9, 0.5, [st], size=17, bold=True, color=AMBER)
    text(s, x, 3.55, 2.9, 0.5, [what], size=20, bold=True, color=WHITE)
    text(s, x, 4.1, 2.8, 1.3, [det], size=15, color=GRAY)
box(s, 0.9, 5.55, 11.5, 1.05, TEAL, radius=0.15)
text(s, 1.2, 5.55, 11.0, 1.05, [[("终点：", {"bold": True, "color": AMBER}),
                                 ("基线核一共实现 27 个 CSR 地址，足够跑起 xv6。", {"color": WHITE})]],
     size=18, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "第 3 步的 12 个是地基。第 4 步加中断的两个；第 6 步加 S 模式，大部分是同一组触发器在 S 地址上的视图；"
         "第 7 步加 satp 开分页。到 xv6 能跑时一共 27 个地址。")

prs.save(OUT)
print("wrote", OUT, "slides:", len(prs.slides))
